"""v138 — Credit File Summary → .xlsx.

The workbook is rendered from a spec the BROWSER sends, so every field is
untrusted input from the shape of the page: sizes are clamped and a value that
Excel would execute as a formula is neutralized. The layout assertions exist
because the owner's requirement was «دقیقاً به همین شکلی که هست» — a workbook
that loses the merges or the band fills is not the form.
"""
import io

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.routers.sheet_export import SheetExportIn, build_sheet_workbook


def _bytes(payload: dict) -> bytes:
    return build_sheet_workbook(SheetExportIn(**payload))


def _wb(payload: dict):
    return load_workbook(io.BytesIO(_bytes(payload)))


def _styles(payload: dict):
    """Resolve each cell's REAL fill/border from the written file.

    openpyxl's reader hands back a bare ``MergedCell`` for every column a merge
    covers, with no fill and no border — even though the file it just wrote
    carries a style index for each of them. Asserting on that read-back would
    have "proved" that the bands lose their colour when Excel shows them
    correctly; so the check reads the XML, which is the ground truth.
    """
    import re
    import zipfile
    from xml.etree import ElementTree as ET

    z = zipfile.ZipFile(io.BytesIO(_bytes(payload)))
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    styles = ET.fromstring(z.read("xl/styles.xml"))
    xfs = styles.find(f"{ns}cellXfs")
    fills = [f.find(f"{ns}patternFill").get("patternType") if f.find(f"{ns}patternFill") is not None else None
             for f in styles.find(f"{ns}fills")]
    fill_rgb = []
    for f in styles.find(f"{ns}fills"):
        pf = f.find(f"{ns}patternFill")
        fg = pf.find(f"{ns}fgColor") if pf is not None else None
        fill_rgb.append((fg.get("rgb") if fg is not None else None))
    borders = [len(list(b)) and any(e.get("style") for e in b) for b in styles.find(f"{ns}borders")]

    sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
    out = {}
    for c in sheet.iter(f"{ns}c"):
        idx = int(c.get("s") or 0)
        xf = list(xfs)[idx]
        fid, bid = int(xf.get("fillId") or 0), int(xf.get("borderId") or 0)
        out[c.get("r")] = {
            "fill": fill_rgb[fid] if fills[fid] == "solid" else None,
            "bordered": bool(borders[bid]),
        }
    return out


BASE_TABLE = {
    "kind": "table",
    "widths": [0.1, 0.3, 0.6],
    "rows": [
        [{"text": "Account Details", "span": 3, "bold": True, "fill": "C7CCD3", "align": "left"}],
        [{"text": "S/No.", "span": 1, "bold": True, "fill": "DDE1E7", "align": "center"},
         {"text": "Description", "span": 1, "bold": True, "fill": "DDE1E7", "align": "center"},
         {"text": "Details", "span": 1, "bold": True, "fill": "DDE1E7", "align": "center"}],
        [{"text": "1", "span": 1, "fill": "EEF1F5", "align": "center"},
         {"text": "Customer's Name", "span": 1, "bold": True, "fill": "EEF1F5"},
         {"text": "Abu Amir Furnishing Branch", "span": 1}],
    ],
}


class TestWorkbookMatchesTheForm:
    def test_it_produces_a_readable_workbook_with_the_values(self):
        ws = _wb({"title": "CREDIT FILE SUMMARY (Corporate)", "account": "351394",
                  "blocks": [BASE_TABLE]}).active
        assert ws["A1"].value == "Account Details"
        assert ws["A3"].value == "1"
        assert ws["C3"].value == "Abu Amir Furnishing Branch"

    def test_a_spanning_cell_becomes_a_real_merge_reaching_the_margin(self):
        wb = _wb({"blocks": [BASE_TABLE]})
        ws = wb.active
        last = get_column_letter(ws.max_column)
        assert f"A1:{last}1" in {str(r) for r in ws.merged_cells.ranges}

    def test_band_and_header_fills_survive_across_the_whole_merge(self):
        st = _styles({"blocks": [BASE_TABLE]})
        # a merged band only LOOKS filled if EVERY covered column carries the
        # fill — the anchor alone leaves the rest of the band white in Excel
        for ref in ("A1", "B1", "C1"):
            assert (st[ref]["fill"] or "").endswith("C7CCD3"), (ref, st[ref])
        assert (st["A2"]["fill"] or "").endswith("DDE1E7")

    def test_every_cell_is_bordered_like_the_printed_grid(self):
        st = _styles({"blocks": [BASE_TABLE]})
        for ref in ("A1", "B1", "C1", "A2", "C3"):
            assert st[ref]["bordered"], (ref, st[ref])

    def test_column_widths_follow_the_forms_proportions(self):
        """The S/No. column must stay narrow and the value column wide, whatever
        the common grid ends up being — measure each cell's TOTAL width, since a
        form column may be made of several grid columns."""
        ws = _wb({"blocks": [BASE_TABLE]}).active
        spans = {rng.min_col: (rng.min_col, rng.max_col)
                 for rng in ws.merged_cells.ranges if rng.min_row == 3}
        def width(col):
            lo, hi = spans.get(col, (col, col))
            return sum(ws.column_dimensions[get_column_letter(c)].width for c in range(lo, hi + 1))
        cols = sorted({c.column for c in ws[3] if c.value})
        a, b, c = (width(x) for x in cols[:3])
        assert a < b < c, (a, b, c)          # 0.1 / 0.3 / 0.6

    def test_the_sheet_prints_on_a4_fitted_to_width_but_never_squashed(self):
        ws = _wb({"blocks": [BASE_TABLE]}).active
        assert ws.page_setup.orientation == "portrait"
        assert ws.page_setup.fitToWidth == 1
        # fitToHeight 0 = "as many pages tall as it needs" — squashing a long
        # form onto one page is what makes it unreadable (the pagination lesson)
        assert ws.page_setup.fitToHeight == 0

    def test_the_date_label_lands_on_a_column_wide_enough_to_read(self):
        """It first rendered in whatever hairline column happened to be second
        from the right, and came out clipped."""
        wb = _wb({"blocks": [
            {"kind": "banner", "left": "BANK SADERAT IRAN", "rightLabel": "Date",
             "rightValue": "26/09/2026"},
            BASE_TABLE,
        ]})
        ws = wb.active
        label = next(c for row in ws.iter_rows() for c in row if c.value == "Date")
        assert ws.column_dimensions[get_column_letter(label.column)].width >= 6, \
            ws.column_dimensions[get_column_letter(label.column)].width

    def test_banner_title_line_and_signatures_all_render(self):
        wb = _wb({"title": "CREDIT FILE SUMMARY (Corporate)", "blocks": [
            {"kind": "banner", "left": "BANK SADERAT IRAN", "leftSub": "U.A.E.",
             "rightLabel": "Date", "rightValue": "26/09/2026"},
            {"kind": "title", "text": "CREDIT FILE SUMMARY (Corporate)"},
            {"kind": "line", "text": "Branch Code and Name: 2690"},
            BASE_TABLE,
            {"kind": "signatures", "left": "Prepared By:", "right": "Authorized:"},
        ]})
        text = " ".join(str(c.value) for row in wb.active.iter_rows() for c in row if c.value)
        for expected in ("BANK SADERAT IRAN", "26/09/2026", "CREDIT FILE SUMMARY (Corporate)",
                         "Branch Code and Name: 2690", "Prepared By:", "Authorized:"):
            assert expected in text, expected

    def test_a_narrow_section_reaches_the_right_margin(self):
        """The first attempt laid every section on the WIDEST table's columns, so
        a 6-column section stopped short of a 7-column sheet and left a bare
        strip down the right-hand side — the owner read it as «یه ستون خالیه».
        Every row must now end at the last column."""
        ws = _wb({"blocks": [
            {"kind": "table", "widths": [0.1, 0.3, 0.6],       # 3 columns
             "rows": [[{"text": "a", "span": 1}, {"text": "b", "span": 1}, {"text": "c", "span": 1}]]},
            {"kind": "table", "widths": [0.25, 0.25, 0.25, 0.25],   # 4 columns
             "rows": [[{"text": "w", "span": 1}, {"text": "x", "span": 1},
                       {"text": "y", "span": 1}, {"text": "z", "span": 1}]]},
        ]}).active
        n = ws.max_column
        for r in (1, 3):
            # the right edge is the furthest of BOTH: a merged range's end, and
            # a plain cell that happens to need no merge
            covered = max(
                [rng.max_col for rng in ws.merged_cells.ranges if rng.min_row == r]
                + [c.column for c in ws[r] if c.value]
            )
            assert covered == n, f"row {r} stops at {covered} of {n}"

    def test_sections_with_different_column_counts_share_one_grid(self):
        """A 6-column and a 7-column section must land on the SAME boundaries,
        or the sheet reads as scattered however neat each table is on its own."""
        ws = _wb({"blocks": [
            {"kind": "table", "widths": [0.5, 0.5],
             "rows": [[{"text": "a", "span": 1}, {"text": "b", "span": 1}]]},
            {"kind": "table", "widths": [0.5, 0.25, 0.25],
             "rows": [[{"text": "c", "span": 1}, {"text": "d", "span": 1}, {"text": "e", "span": 1}]]},
        ]}).active
        def last_col_of_first_cell(row: int) -> int:
            for rng in ws.merged_cells.ranges:
                if rng.min_row == row and rng.min_col == 1:
                    return rng.max_col
            return 1
        # "a" and "c" each take the first half, so they must end on the SAME
        # boundary — that is what "one grid" means visually
        assert last_col_of_first_cell(1) == last_col_of_first_cell(3)

    def test_the_grid_never_grows_into_hairline_columns(self):
        """Boundaries that nearly agree are snapped together; an unbounded union
        would produce a sheet of 2-character columns that nobody can edit."""
        blocks = [{"kind": "table", "widths": [0.2 + i * 0.001, 0.8 - i * 0.001],
                   "rows": [[{"text": "x", "span": 1}, {"text": "y", "span": 1}]]}
                  for i in range(10)]
        ws = _wb({"blocks": blocks}).active
        assert ws.max_column <= 12, ws.max_column


class TestUntrustedInput:
    @pytest.mark.parametrize("bad", ["=1+1", "+1", "-1", "@SUM(A1)",
                                     '=HYPERLINK("http://x","click")'])
    def test_a_value_excel_would_execute_is_neutralized(self, bad):
        ws = _wb({"blocks": [{"kind": "table", "widths": [1.0],
                              "rows": [[{"text": bad, "span": 1}]]}]}).active
        assert ws["A1"].value.startswith("'"), ws["A1"].value

    def test_an_ordinary_value_is_left_exactly_as_it_is(self):
        ws = _wb({"blocks": [{"kind": "table", "widths": [1.0],
                              "rows": [[{"text": "1,000,000", "span": 1}]]}]}).active
        assert ws["A1"].value == "1,000,000"

    def test_a_span_wider_than_the_grid_is_clamped_instead_of_crashing(self):
        ws = _wb({"blocks": [{"kind": "table", "widths": [0.5, 0.5],
                              "rows": [[{"text": "wide", "span": 20}]]}]}).active
        assert ws["A1"].value == "wide"
        assert all(str(r).endswith("1") for r in ws.merged_cells.ranges)

    def test_an_empty_or_absent_widths_list_still_renders(self):
        ws = _wb({"blocks": [{"kind": "table", "widths": [],
                              "rows": [[{"text": "x", "span": 1}]]}]}).active
        assert ws["A1"].value == "x"

    def test_an_unknown_block_kind_is_ignored_not_fatal(self):
        ws = _wb({"blocks": [{"kind": "mystery", "text": "?"}, BASE_TABLE]}).active
        assert ws["A1"].value == "Account Details"

    def test_persian_text_round_trips(self):
        ws = _wb({"blocks": [{"kind": "table", "widths": [1.0],
                              "rows": [[{"text": "املاک / Mortgaged Properties", "span": 1}]]}]}).active
        assert ws["A1"].value == "املاک / Mortgaged Properties"


class TestEndpoint:
    async def test_it_returns_an_xlsx_named_after_the_account(self, client, auth_headers):
        r = await client.post("/api/credit-file/export-excel", headers=auth_headers, json={
            "title": "CREDIT FILE SUMMARY (Corporate)", "account": "351394",
            "blocks": [BASE_TABLE],
        })
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert "CreditFile-Corporate-351394.xlsx" in r.headers["content-disposition"]
        assert load_workbook(io.BytesIO(r.content)).active["A1"].value == "Account Details"

    async def test_an_account_number_cannot_escape_the_file_name(self, client, auth_headers):
        r = await client.post("/api/credit-file/export-excel", headers=auth_headers, json={
            "title": "Retail", "account": "../../etc/passwd", "blocks": [BASE_TABLE],
        })
        assert r.status_code == 200, r.text
        assert "/" not in r.headers["content-disposition"].split("filename=")[1]

    async def test_a_spec_with_no_blocks_is_rejected(self, client, auth_headers):
        r = await client.post("/api/credit-file/export-excel", headers=auth_headers,
                              json={"title": "x", "account": "1", "blocks": []})
        assert r.status_code == 422

    async def test_it_requires_a_logged_in_user(self, client):
        r = await client.post("/api/credit-file/export-excel",
                              json={"title": "x", "account": "1", "blocks": [BASE_TABLE]})
        assert r.status_code in (401, 403), r.status_code
