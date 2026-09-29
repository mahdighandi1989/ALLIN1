"""v146 — samples attached to an inspection sheet, and the duty to read them.

The owner asked for uploads of ANY type up to 100 MB, then added the sentence
that decides the whole design: «حجم هم باعث نشه ناظر نتونه بگه من نمیخونمش» —
size must never become an excuse not to read it.

So these tests are mostly about that: the text is extracted once at upload, the
reviewer reads slices, the slices are counted, and an answer is REFUSED while a
sample is unread. Everything else (where the bytes live, what happens when Drive
is missing, what «no text» means) is tested because getting those wrong is how a
sample silently disappears or silently looks empty.
"""
import io
import json

import pytest

from app.models.inspection import (
    STATUS_ANSWERED, STATUS_OPEN, InspectionFile, file_read_debt,
)
from app.services import inspection_files as ifiles

SPOT = {
    "page": "/settings", "page_label": "تنظیمات",
    "section_id": "formats", "section_label": "قالبِ اسناد",
    "reopen": "/settings#formats",
    "dom_path": "section#formats", "covered_text": "قالبِ سند",
    "rect": {"x": 1, "y": 2, "w": 3, "h": 4}, "viewport": {"w": 1200, "h": 800},
}


@pytest.fixture(autouse=True)
def _drive_off(monkeypatch, tmp_path):
    """No Drive in tests — the LOCAL fallback path is what runs here, and it must
    work AND announce itself as non-durable."""
    monkeypatch.setenv("INSPECTION_FILE_DIR", str(tmp_path / "insp"))
    monkeypatch.setattr("app.services.google_drive.is_configured", lambda: False)


async def _sheet(client, headers, text="قالبِ سندِ جدید می‌خواهم"):
    r = await client.post("/api/inspection", headers=headers,
                          json={"text": text, "spot": SPOT})
    assert r.status_code == 200, r.text
    return r.json()["report"]


async def _upload(client, headers, report_id, name, data, caption="", mime=None):
    files = {"file": (name, io.BytesIO(data), mime or "application/octet-stream")}
    r = await client.post(f"/api/inspection/{report_id}/files", headers=headers,
                          files=files, data={"caption": caption})
    return r


# ---------------------------------------------------------------------------
class TestAnyFileType:
    async def test_a_plain_text_sample_is_stored_and_its_text_extracted(
            self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        r = await _upload(client, auth_headers, rep["id"], "نمونه.txt",
                          "سلام\nاین یک نمونه است".encode(),
                          caption="این ساختارِ موردِ نظرِ من است")
        assert r.status_code == 200, r.text
        f = r.json()["file"]
        assert f["filename"] == "نمونه.txt"
        assert f["extract_status"] == "ok"
        assert f["text_chars"] > 0
        # the owner's words travel WITH the bytes — «هم به توضیحات»
        assert f["caption"] == "این ساختارِ موردِ نظرِ من است"
        assert f["read_chars"] == 0 and f["fully_read"] is False

    async def test_a_non_image_binary_is_accepted_and_says_why_it_has_no_text(
            self, client, auth_headers):
        """«unsupported» must never be reported as «empty»: we did not try, and
        the reviewer has to open the file itself."""
        rep = await _sheet(client, auth_headers)
        r = await _upload(client, auth_headers, rep["id"], "sample.bin", b"\x00\x01\x02\x03")
        f = r.json()["file"]
        assert f["extract_status"] == "unsupported"
        assert f["extract_note"]                      # the reason is mandatory
        assert "خودِ فایل" in f["extract_note"]

    async def test_an_image_from_outside_is_an_image_not_an_empty_file(
            self, client, auth_headers):
        """The owner's case: «عکسی که از بیرون دانلود کردم به عنوان نمونه»."""
        rep = await _sheet(client, auth_headers)
        r = await _upload(client, auth_headers, rep["id"], "ref.png", b"\x89PNG\r\n\x1a\n",
                          mime="image/png")
        f = r.json()["file"]
        assert f["extract_status"] == "image"
        assert f["text_chars"] == 0

    async def test_a_word_sample_of_a_document_format_is_read_including_tables(
            self, client, auth_headers):
        """«یه فرمت سند اضافه کنم و نسخه ورد نمونه بهش بدم تا ایجاد کنه» — the
        content of a form sample lives in its TABLES, so they must come through."""
        from docx import Document
        doc = Document()
        doc.add_paragraph("قالبِ نامهٔ پیشنهاد")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "نامِ مشتری"
        t.cell(0, 1).text = "مبلغ"
        t.cell(1, 0).text = "شرکتِ الف"
        t.cell(1, 1).text = "۱۲۳"
        buf = io.BytesIO()
        doc.save(buf)

        rep = await _sheet(client, auth_headers)
        r = await _upload(
            client, auth_headers, rep["id"], "قالب.docx", buf.getvalue(),
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert r.status_code == 200, r.text
        fid = r.json()["file"]["id"]
        got = (await client.get(f"/api/inspection/files/{fid}/text",
                                headers=auth_headers)).json()
        assert "قالبِ نامهٔ پیشنهاد" in got["text"]
        assert "نامِ مشتری" in got["text"] and "شرکتِ الف" in got["text"]

    async def test_a_pdf_sample_carries_its_page_numbers(self, client, auth_headers):
        pytest.importorskip("pypdf")
        from pypdf import PdfWriter
        w = PdfWriter()
        w.add_blank_page(width=200, height=200)
        w.add_blank_page(width=200, height=200)
        buf = io.BytesIO()
        w.write(buf)
        rep = await _sheet(client, auth_headers)
        r = await _upload(client, auth_headers, rep["id"], "s.pdf", buf.getvalue(),
                          mime="application/pdf")
        f = r.json()["file"]
        # blank pages have no text layer — that is «unsupported», with the count,
        # NOT «empty»: a scanned PDF must send the reviewer to the file itself
        assert f["extract_status"] == "unsupported"
        assert f["page_count"] == 2
        assert "اسکن" in f["extract_note"]

    async def test_a_spreadsheet_sample_is_flattened_to_text(self, client, auth_headers):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active["A1"] = "ستون"
        wb.active["B1"] = 42
        buf = io.BytesIO()
        wb.save(buf)
        rep = await _sheet(client, auth_headers)
        r = await _upload(client, auth_headers, rep["id"], "n.xlsx", buf.getvalue())
        assert r.json()["file"]["extract_status"] == "ok"


class TestPdfTextLayerDecision:
    """`ok` vs `unsupported` for a PDF is decided by how many pages carried real
    text — never by whether the assembled string is non-blank.

    The page markers the extractor inserts («--- صفحهٔ ۳ ---») are themselves
    text. Before this was fixed, a scanned PDF came back «readable», and a
    reviewer could clear its whole read debt on forty characters of markers.
    """

    def test_no_page_with_text_is_unsupported_not_ok_and_not_empty(self, monkeypatch):
        monkeypatch.setattr(ifiles, "_pdf_text",
                            lambda data: ("\n--- صفحهٔ 1 ---\n\n--- صفحهٔ 2 ---\n", 2, 0))
        out = ifiles.extract(b"%PDF", "s.pdf", "application/pdf")
        assert out["status"] == "unsupported"
        assert out["text"] == ""            # nothing pretending to be content
        assert out["page_count"] == 2
        assert "اسکن" in out["note"]

    def test_every_page_with_text_is_ok(self, monkeypatch):
        monkeypatch.setattr(ifiles, "_pdf_text",
                            lambda data: ("\n--- صفحهٔ 1 ---\nقرارداد", 1, 1))
        out = ifiles.extract(b"%PDF", "s.pdf", "application/pdf")
        assert out["status"] == "ok" and "قرارداد" in out["text"]
        assert out["note"] == "از 1 صفحهٔ PDF"

    def test_a_partly_scanned_pdf_is_readable_but_says_what_is_missing(self, monkeypatch):
        """Half a document is worse than none if nobody is told which half."""
        monkeypatch.setattr(ifiles, "_pdf_text",
                            lambda data: ("\n--- صفحهٔ 1 ---\nمتن", 10, 3))
        out = ifiles.extract(b"%PDF", "s.pdf", "application/pdf")
        assert out["status"] == "ok"
        assert "فقط 3 صفحه" in out["note"] and "اسکن" in out["note"]

    def test_a_corrupt_pdf_is_failed_with_its_reason_not_a_crash(self):
        out = ifiles.extract(b"%PDF-1.4 truncated", "broken.pdf", "application/pdf")
        assert out["status"] == "failed"
        assert out["note"] and "استخراج شکست خورد" in out["note"]

    def test_the_text_ceiling_is_reported_not_silently_applied(self, monkeypatch):
        monkeypatch.setattr(ifiles, "MAX_TEXT_CHARS", 50)
        out = ifiles.extract(("ز" * 500).encode(), "long.txt")
        assert out["status"] == "ok"
        assert len(out["text"]) == 50
        assert "بریده شد" in out["note"]
        # the flag is what stops «read the part we kept» counting as «read it»
        assert out["truncated"] is True

    def test_an_untruncated_file_is_not_flagged(self):
        assert ifiles.extract(b"short", "s.txt")["truncated"] is False

    def test_the_default_ceiling_is_big_enough_for_a_real_document(self):
        """4M characters made a 60MB sample a 6% read. A real document sample —
        even a 6,000-page one — never reaches 20M."""
        assert ifiles.MAX_TEXT_CHARS >= 20_000_000


class TestATruncatedSampleIsNotFullyRead:
    """The trap this closes: text cut at a ceiling, `fully_read` true anyway.

    That is a cap reported as a total — the same fault v144 had to fix in the
    data-quality reports. Here it would have let a reviewer discharge its duty on
    a fraction of a 60MB sample and believe it had read the whole thing.
    """

    @pytest.fixture(autouse=True)
    def _as_supervisor(self, monkeypatch, test_user):
        monkeypatch.setenv("SUPERVISOR_API_USER", test_user.username)

    async def test_finishing_a_truncated_text_is_not_enough(
            self, client, auth_headers, monkeypatch):
        monkeypatch.setattr(ifiles, "MAX_TEXT_CHARS", 100)
        monkeypatch.setattr(ifiles, "SLICE_CHARS", 100)
        rep = await _sheet(client, auth_headers)
        f = (await _upload(client, auth_headers, rep["id"], "huge.txt",
                           ("و" * 900).encode())).json()["file"]
        assert f["text_truncated"] is True

        # read every character we kept …
        got = (await client.get(f"/api/inspection/files/{f['id']}/text",
                                headers=auth_headers)).json()
        assert got["fully_read"] is True and got["has_more"] is False
        # … and the sheet is STILL not answerable, because that was not the file
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "خواندم", "outcome": "not-done"})
        assert r.status_code == 422, r.text
        assert "بریده" in r.json()["detail"]

        # opening the file itself is what completes it
        assert (await client.get(f"/api/inspection/files/{f['id']}/raw",
                                 headers=auth_headers)).status_code == 200
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "خواندم و خودِ فایل را هم دیدم", "outcome": "not-done"})
        assert r.status_code == 200, r.text

    async def test_the_debt_says_WHY_it_is_a_debt(self):
        class F:
            id, filename, extract_status = "1", "big.txt", "ok"
            text_chars, read_chars = 100, 100
            text_truncated, viewed_at = True, None
        d = file_read_debt([F()])
        assert d and d[0]["reason"] == "truncated"

    async def test_an_untruncated_fully_read_file_is_clear(self):
        class F:
            id, filename, extract_status = "1", "ok.txt", "ok"
            text_chars, read_chars = 100, 100
            text_truncated, viewed_at = False, None
        assert file_read_debt([F()]) == []


class TestSizeIsNotAnExcuse:
    async def test_the_text_is_served_in_slices_with_its_true_total(
            self, client, auth_headers, monkeypatch):
        monkeypatch.setattr(ifiles, "SLICE_CHARS", 100)
        rep = await _sheet(client, auth_headers)
        body = ("ا" * 250).encode()
        fid = (await _upload(client, auth_headers, rep["id"], "long.txt", body)).json()["file"]["id"]

        seen = ""
        offset, guard = 0, 0
        while True:
            got = (await client.get(f"/api/inspection/files/{fid}/text?offset={offset}",
                                    headers=auth_headers)).json()
            # never report a slice as the whole — the v144 discipline again
            assert got["text_chars"] == 250
            assert got["offset"] == offset
            seen += got["text"]
            if not got["has_more"]:
                break
            offset = got["next_offset"]
            guard += 1
            assert guard < 10
        assert len(seen) == 250
        assert got["fully_read"] is True
        assert got["read_chars"] == 250

    async def test_reading_only_the_last_slice_does_not_clear_the_debt(
            self, client, auth_headers, monkeypatch):
        """Otherwise one call to the end would count as having read everything —
        exactly the loophole this feature exists to close."""
        monkeypatch.setattr(ifiles, "SLICE_CHARS", 50)
        rep = await _sheet(client, auth_headers)
        fid = (await _upload(client, auth_headers, rep["id"], "l.txt",
                             ("ب" * 200).encode())).json()["file"]["id"]
        got = (await client.get(f"/api/inspection/files/{fid}/text?offset=150",
                                headers=auth_headers)).json()
        assert got["read_chars"] == 0          # the gap is still unread
        assert got["fully_read"] is False

    async def test_a_file_over_the_ceiling_is_refused_with_the_ceiling_named(
            self, client, auth_headers, monkeypatch):
        monkeypatch.setattr(ifiles, "MAX_BYTES", 1024)
        monkeypatch.setattr(ifiles, "MAX_MB", 1)
        rep = await _sheet(client, auth_headers)
        r = await _upload(client, auth_headers, rep["id"], "big.txt", b"x" * 5000)
        assert r.status_code == 413
        assert "مگابایت" in r.json()["detail"]
        assert "INSPECTION_MAX_FILE_MB" in r.json()["detail"]

    async def test_the_default_ceiling_is_the_hundred_megabytes_asked_for(self):
        assert ifiles.MAX_MB >= 50
        assert ifiles.MAX_BYTES == ifiles.MAX_MB * 1024 * 1024

    async def test_an_empty_upload_is_refused(self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        r = await _upload(client, auth_headers, rep["id"], "e.txt", b"")
        assert r.status_code == 422


class TestTheSupervisorMustRead:
    @pytest.fixture(autouse=True)
    def _as_supervisor(self, monkeypatch, test_user):
        monkeypatch.setenv("SUPERVISOR_API_USER", test_user.username)

    async def test_an_answer_is_refused_while_a_sample_is_unread(
            self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        await _upload(client, auth_headers, rep["id"], "spec.txt",
                      ("قاعده " * 200).encode(), caption="این را بخوان")
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "انجام شد", "outcome": "not-done"})
        assert r.status_code == 422, r.text
        d = r.json()["detail"]
        assert "spec.txt" in d and "نویسه" in d
        # the way to comply is IN the refusal
        assert "/text" in d

    async def test_once_read_to_the_end_the_answer_goes_through(
            self, client, auth_headers, monkeypatch):
        monkeypatch.setattr(ifiles, "SLICE_CHARS", 100)
        rep = await _sheet(client, auth_headers)
        fid = (await _upload(client, auth_headers, rep["id"], "spec.txt",
                             ("ج" * 320).encode())).json()["file"]["id"]
        offset = 0
        while True:
            got = (await client.get(f"/api/inspection/files/{fid}/text?offset={offset}",
                                    headers=auth_headers)).json()
            if not got["has_more"]:
                break
            offset = got["next_offset"]
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "خواندم و انجام شد", "outcome": "not-done"})
        assert r.status_code == 200, r.text
        assert r.json()["report"]["status"] == STATUS_ANSWERED

    async def test_an_image_sample_must_be_OPENED_not_waved_through(
            self, client, auth_headers):
        """An image has no text, so the text guard cannot see it. Looking at it
        IS fetching it — so that is what gets demanded and recorded."""
        rep = await _sheet(client, auth_headers)
        fid = (await _upload(client, auth_headers, rep["id"], "ref.png",
                             b"\x89PNG\r\n\x1a\n", mime="image/png")).json()["file"]["id"]
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "دیدم", "outcome": "not-done"})
        assert r.status_code == 422
        assert "باز نشده" in r.json()["detail"]

        assert (await client.get(f"/api/inspection/files/{fid}/raw",
                                 headers=auth_headers)).status_code == 200
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "دیدم", "outcome": "not-done"})
        assert r.status_code == 200, r.text

    async def test_an_unreadable_file_must_also_be_opened(self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        fid = (await _upload(client, auth_headers, rep["id"], "x.bin",
                             b"\x00\x01")).json()["file"]["id"]
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "?", "outcome": "not-done"})
        assert r.status_code == 422
        await client.get(f"/api/inspection/files/{fid}/raw", headers=auth_headers)
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "?", "outcome": "not-done"})
        assert r.status_code == 200

    async def test_the_supervisor_cannot_delete_the_sample(self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        fid = (await _upload(client, auth_headers, rep["id"], "a.txt", b"hi")).json()["file"]["id"]
        r = await client.delete(f"/api/inspection/files/{fid}", headers=auth_headers)
        assert r.status_code == 403

    async def test_the_queue_shows_the_reading_it_owes_before_it_answers(
            self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        await _upload(client, auth_headers, rep["id"], "q.txt", ("د" * 90).encode())
        q = (await client.get("/api/inspection/queue", headers=auth_headers)).json()
        assert q["files_to_read"] == 1
        sheet = next(x for x in q["reports"] if x["id"] == rep["id"])
        assert sheet["files"][0]["filename"] == "q.txt"
        assert sheet["read_debt"][0]["remaining"] == 90
        # the listing must NOT ship the extracted text
        assert "text" not in sheet["files"][0]


class TestTheOwnerSide:
    async def test_the_owner_is_never_blocked_by_the_read_guard(
            self, client, auth_headers):
        """The duty is the supervisor's. The owner uploading and then writing
        more must not be refused because of their own file."""
        rep = await _sheet(client, auth_headers)
        await _upload(client, auth_headers, rep["id"], "mine.txt", ("ه" * 500).encode())
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "این هم توضیحِ بیشتر"})
        assert r.status_code == 200, r.text

    async def test_the_owner_can_remove_a_sample(self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        fid = (await _upload(client, auth_headers, rep["id"], "a.txt", b"hi")).json()["file"]["id"]
        assert (await client.delete(f"/api/inspection/files/{fid}",
                                    headers=auth_headers)).status_code == 200
        assert (await client.get(f"/api/inspection/files/{fid}",
                                 headers=auth_headers)).status_code == 404

    async def test_a_new_sample_reopens_an_answered_sheet(self, client, auth_headers,
                                                          monkeypatch, test_user):
        """New material the previous answer never saw puts the sheet back in the
        queue — the same rule an owner follow-up note already followed."""
        monkeypatch.setenv("SUPERVISOR_API_USER", test_user.username)
        rep = await _sheet(client, auth_headers)
        r = await client.post(f"/api/inspection/{rep['id']}/notes", headers=auth_headers,
                              json={"text": "بررسی شد", "outcome": "not-done"})
        assert r.json()["report"]["status"] == STATUS_ANSWERED
        monkeypatch.delenv("SUPERVISOR_API_USER")
        r = await _upload(client, auth_headers, rep["id"], "more.txt", b"more")
        assert r.json()["report"]["status"] == STATUS_OPEN


class TestWhereTheBytesLive:
    async def test_without_drive_the_file_still_works_but_says_it_is_not_durable(
            self, client, auth_headers):
        """Never lose the capability — but never let a fallback look permanent.
        The container disk is wiped on every deploy (OPEN_ITEMS #2)."""
        rep = await _sheet(client, auth_headers)
        f = (await _upload(client, auth_headers, rep["id"], "a.txt", b"hi")).json()["file"]
        assert f["store"] == "local"
        assert f["durable"] is False
        assert "درایو" in f["store_note"] and "دیپلوی" in f["store_note"]

    async def test_with_drive_configured_the_file_goes_to_its_own_folder(
            self, client, auth_headers, monkeypatch):
        seen = {}

        def fake_upload(**kw):
            seen.update(kw)
            return {"id": "drive123", "link": "https://drive.example/f/drive123",
                    "folder_id": "fold", "action": "created"}

        monkeypatch.setattr("app.services.google_drive.is_configured", lambda: True)
        monkeypatch.setattr("app.services.google_drive.upload_file", fake_upload)
        rep = await _sheet(client, auth_headers)
        f = (await _upload(client, auth_headers, rep["id"], "spec.txt", b"hi")).json()["file"]
        assert f["store"] == "drive" and f["durable"] is True
        assert f["drive_link"] == "https://drive.example/f/drive123"
        # its own folder, per sheet, created on demand
        assert seen["path_parts"] == ["inspection", f"report-{rep['number']}"]

    async def test_a_drive_failure_falls_back_and_records_the_reason(
            self, client, auth_headers, monkeypatch):
        def boom(**kw):
            raise RuntimeError("no permission")

        monkeypatch.setattr("app.services.google_drive.is_configured", lambda: True)
        monkeypatch.setattr("app.services.google_drive.upload_file", boom)
        rep = await _sheet(client, auth_headers)
        f = (await _upload(client, auth_headers, rep["id"], "a.txt", b"hi")).json()["file"]
        assert f["store"] == "local"
        assert "no permission" in f["store_note"]

    async def test_the_raw_bytes_come_back_unchanged(self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        payload = "متنِ نمونه".encode()
        fid = (await _upload(client, auth_headers, rep["id"], "a.txt", payload)).json()["file"]["id"]
        r = await client.get(f"/api/inspection/files/{fid}/raw", headers=auth_headers)
        assert r.status_code == 200
        assert r.content == payload

    async def test_a_filename_cannot_escape_its_folder(self):
        assert "/" not in ifiles.safe_filename("../../etc/passwd")
        assert ifiles.safe_filename("../../etc/نمونه.docx") == "نمونه.docx"


class TestDeletingASheetTakesItsSamples:
    """v149 — the missing link the dependency walk found.

    `InspectionFile` arrived in v146 and `delete_report` was never updated, so a
    deleted sheet left its file rows behind. Every listing filters by
    `report_id`, so they were reachable from nowhere in the product while still
    holding up to 20M characters each — the orphaned-import-attachment finding
    again: not lost, but not findable, which for whoever is looking is the same.
    The neighbour gave it away: shots were already cleaned up two lines above.
    """

    async def test_the_file_rows_go_with_the_sheet(self, client, auth_headers, db_session):
        from app.models.inspection import InspectionFile
        from sqlalchemy import select

        rep = await _sheet(client, auth_headers)
        fid = (await _upload(client, auth_headers, rep["id"], "spec.txt",
                             b"content")).json()["file"]["id"]
        assert (await client.delete(f"/api/inspection/{rep['id']}",
                                    headers=auth_headers)).status_code == 200

        left = (await db_session.execute(select(InspectionFile).where(
            InspectionFile.report_id == rep["id"]))).scalars().all()
        assert left == [], "a deleted sheet must not leave its samples behind"
        assert (await client.get(f"/api/inspection/files/{fid}",
                                 headers=auth_headers)).status_code == 404

    async def test_it_reports_how_many_it_removed(self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        await _upload(client, auth_headers, rep["id"], "a.txt", b"a")
        await _upload(client, auth_headers, rep["id"], "b.txt", b"b")
        r = await client.delete(f"/api/inspection/{rep['id']}", headers=auth_headers)
        assert r.json()["files_removed"] == 2

    async def test_a_sheet_with_no_samples_still_deletes_cleanly(
            self, client, auth_headers):
        rep = await _sheet(client, auth_headers)
        r = await client.delete(f"/api/inspection/{rep['id']}", headers=auth_headers)
        assert r.status_code == 200 and r.json()["files_removed"] == 0

    async def test_the_drive_copy_is_left_alone(self, client, auth_headers, monkeypatch):
        """Rule 2 is quarantine, not deletion — and a sample the owner sent is
        evidence. Deleting the row must never reach for the bytes."""
        deleted: list = []
        monkeypatch.setattr("app.services.google_drive.is_configured", lambda: True)
        monkeypatch.setattr("app.services.google_drive.upload_file",
                            lambda **kw: {"id": "d1", "link": "http://x/d1",
                                          "folder_id": "f", "action": "created"})
        monkeypatch.setattr("app.services.google_drive.delete_file",
                            lambda fid: deleted.append(fid))
        rep = await _sheet(client, auth_headers)
        await _upload(client, auth_headers, rep["id"], "a.txt", b"a")
        await client.delete(f"/api/inspection/{rep['id']}", headers=auth_headers)
        assert deleted == [], "the Drive copy must survive the sheet"


class TestReadDebtIsHonest:
    def test_a_partially_read_file_is_still_a_debt(self):
        class F:
            id, filename, extract_status = "1", "a.txt", "ok"
            text_chars, read_chars, viewed_at = 100, 40, None
        d = file_read_debt([F()])
        assert d and d[0]["remaining"] == 60

    def test_a_fully_read_file_is_no_debt(self):
        class F:
            id, filename, extract_status = "1", "a.txt", "ok"
            text_chars, read_chars, viewed_at = 100, 100, None
        assert file_read_debt([F()]) == []

    def test_an_empty_file_is_no_debt_because_there_is_nothing_to_read(self):
        class F:
            id, filename, extract_status = "1", "a.txt", "empty"
            text_chars, read_chars, viewed_at = 0, 0, None
        assert file_read_debt([F()]) == []

    def test_a_failed_extraction_is_no_read_debt(self):
        """It is not the reviewer's fault the extractor broke, and demanding a
        read it cannot perform would deadlock the sheet."""
        class F:
            id, filename, extract_status = "1", "a.bin", "failed"
            text_chars, read_chars, viewed_at = 0, 0, None
        assert file_read_debt([F()]) == []

    def test_no_files_is_no_debt(self):
        assert file_read_debt([]) == [] and file_read_debt(None) == []
