"""v138 — Credit File Summary → a real .xlsx laid out like the printed form.

The client reads its own rendered sheet into a neutral spec (blocks of rows,
merged cells, column proportions, fills) and posts it here; this module renders
it with openpyxl. That split follows the pattern the voucher export established
(v105) and the artifact philosophy recorded in `experiences/`: **the client
describes, the server renders with a real library** — deterministic, validated,
no side-channel storage, the response IS the file.

Why the spec comes from the browser rather than being rebuilt from the database:
the form on screen carries the officer's UNSAVED edits and their «حذف از پرینت»
choices. Rebuilding it server-side would quietly export a different document
from the one they are looking at, which is the one thing the request asked for
(«دقیقاً به همین شکلی که هست»).

Everything that arrives is untrusted: sizes are clamped, text is truncated, and
a value that could be read as a formula is neutralized (Excel executes a cell
starting with = + - @, and a customer name really can begin with one).
"""
from __future__ import annotations

import io
import re
from copy import copy
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field

from app.routers.auth import get_current_active_user

router = APIRouter(tags=["sheet-export"], dependencies=[Depends(get_current_active_user)])

MAX_ROWS = 400
MAX_COLS = 24
MAX_BLOCKS = 40
MAX_TEXT = 2000

_HEX = re.compile(r"^[0-9A-Fa-f]{6}$")
# Arial 10 in Excel: one column-width unit ≈ 1.87 mm. The printed sheet is 186 mm
# wide, so the whole grid gets this many units to share out.
_TOTAL_WIDTH_UNITS = 99.0


class Cell(BaseModel):
    text: str = Field("", max_length=MAX_TEXT)
    span: int = Field(1, ge=1, le=MAX_COLS)
    rowSpan: Optional[int] = Field(None, ge=1, le=50)
    bold: Optional[bool] = None
    align: Optional[str] = Field(None, max_length=6)
    fill: Optional[str] = Field(None, max_length=6)
    wrap: Optional[bool] = None


class Block(BaseModel):
    kind: str = Field(..., max_length=12)
    # table
    rows: List[List[Cell]] = Field(default_factory=list, max_length=MAX_ROWS)
    widths: List[float] = Field(default_factory=list, max_length=MAX_COLS)
    # banner
    left: str = Field("", max_length=300)
    leftSub: str = Field("", max_length=300)
    rightLabel: str = Field("", max_length=60)
    rightValue: str = Field("", max_length=60)
    # title / line
    text: str = Field("", max_length=600)
    # signatures
    right: str = Field("", max_length=120)


class SheetExportIn(BaseModel):
    title: str = Field("CREDIT FILE SUMMARY", max_length=200)
    account: str = Field("", max_length=60)
    blocks: List[Block] = Field(..., min_length=1, max_length=MAX_BLOCKS)


def _safe(v: str) -> str:
    """Neutralize spreadsheet formula injection (same rule as the CSV exporter).

    Excel executes a cell whose text starts with = + - @. A customer name or a
    remark really can start with one, and the officer opening the file would be
    running it.
    """
    s = (v or "")[:MAX_TEXT]
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s


def _fill(v: Optional[str]):
    from openpyxl.styles import PatternFill
    return PatternFill("solid", fgColor=v.upper()) if v and _HEX.match(v) else None



def _norm(widths: List[float]) -> List[float]:
    """Column proportions → cumulative boundaries [0.0 … 1.0]."""
    w = [max(0.0, float(x)) for x in widths]
    tot = sum(w)
    if tot <= 0:
        n = max(1, len(w))
        w, tot = [1.0] * n, float(n)
    out, run = [0.0], 0.0
    for x in w:
        run += x / tot
        out.append(run)
    out[-1] = 1.0
    return out


# Two boundaries closer than this are the same line. Without a tolerance, tables
# whose columns very nearly agree would each contribute their own boundary and
# the grid would fill up with hairline columns.
# Column proportions of the banner row (name | «Date» | the date itself). They
# are declared once and fed to BOTH the grid builder and the renderer, so the
# label always lands on a real boundary wide enough to hold it.
_BANNER = (0.70, 0.10, 0.20)

# Two boundaries closer than this are the same line. Too small and the grid
# fills with hairline columns nobody can edit; too large and a section's
# proportions visibly drift from the form.
_SNAP = 0.028
_MAX_GRID = 12


def _common_grid(blocks) -> List[float]:
    """The union of every table's column boundaries, snapped and capped.

    This is what lets a 6-column section and a 7-column section share one
    worksheet without either of them going ragged.
    """
    marks = {0.0, 1.0}
    # the banner's own splits — ONLY when there is a banner, or a document
    # without one pays for two boundaries it never uses (they showed up as an
    # extra pair of columns splitting a wide value cell for no reason)
    if any(b.kind == "banner" for b in blocks):
        marks.update(_norm(list(_BANNER)))
    for b in blocks:
        if b.kind == "table" and b.widths:
            marks.update(_norm(list(b.widths)[:MAX_COLS]))

    grid: List[float] = []
    for m in sorted(marks):
        if not grid or m - grid[-1] >= _SNAP:
            grid.append(round(m, 4))
        # else: snap onto the boundary already there
    if grid[-1] < 1.0:
        grid[-1] = 1.0
    # cap the width of the grid by dropping the narrowest interior boundaries
    while len(grid) - 1 > _MAX_GRID:
        gaps = [(grid[i + 1] - grid[i], i) for i in range(1, len(grid) - 1)]
        grid.pop(min(gaps)[1])
    return grid


def _cols_for(grid: List[float], bounds: List[float], c0: int, span: int) -> tuple:
    """Map one HTML cell onto the common grid → (first column, last column), 1-based."""
    lo, hi = bounds[c0], bounds[min(c0 + span, len(bounds) - 1)]
    i0 = min(range(len(grid)), key=lambda i: abs(grid[i] - lo))
    i1 = min(range(len(grid)), key=lambda i: abs(grid[i] - hi))
    if i1 <= i0:
        i1 = min(i0 + 1, len(grid) - 1)
    return i0 + 1, i1          # grid index i spans worksheet column i+1


def _extend(ws, row: int, last_col: int, ncols: int, border) -> None:
    """Stretch a row's LAST cell out to the right margin.

    A section narrower than the grid would otherwise stop mid-sheet and leave a
    bare vertical strip down the right-hand side — which is exactly what read as
    "a spare empty column". Parking empty bordered cells there looks no better;
    the last real cell simply has to reach the edge.
    """
    anchor = last_col
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row == row and rng.max_row == row and rng.max_col == last_col:
            anchor = rng.min_col
            ws.unmerge_cells(start_row=row, start_column=anchor, end_row=row, end_column=last_col)
            break
    src = ws.cell(row=row, column=anchor)
    fill = src.fill if src.fill and src.fill.patternType else None
    ws.merge_cells(start_row=row, start_column=anchor, end_row=row, end_column=ncols)
    for c in range(anchor, ncols + 1):
        tgt = ws.cell(row=row, column=c)
        tgt.border = border
        if fill is not None:
            tgt.fill = copy(fill)


def build_sheet_workbook(payload: SheetExportIn) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    # ---- ONE COMMON COLUMN GRID ------------------------------------------
    # A worksheet has a single column grid, but the form is seven tables with
    # DIFFERENT column counts (6 / 6 / 5 / 6 / 7 / 6 / 2). Laying them all on
    # the widest table's columns — the first attempt — left column G empty for
    # every 6-column section, ended each band ragged short of the right edge,
    # and crushed the 2-column Status table into the narrow S/No. column. It
    # read as scattered, because it was.
    #
    # Instead: take every table's column BOUNDARIES as fractions of the width,
    # union them, and lay all of them on the resulting grid. Then each section
    # keeps its own proportions, every boundary is real, every row reaches the
    # right edge, and nothing is left over.
    grid = _common_grid(payload.blocks)
    ncols = len(grid) - 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Credit File"

    thin = Side(style="thin", color="000000")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    def put(r: int, c1: int, c2: int, text: str, *, bold=False, align="left",
            fill: Optional[str] = None, wrap=False, size=8):
        c2 = max(c1, min(c2, ncols))
        cell = ws.cell(row=r, column=c1, value=_safe(text))
        cell.font = Font(name="Arial", size=size, bold=bold)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        # ORDER MATTERS, and it is the opposite of what it looks like:
        # openpyxl's merge_cells CLEARS the fill of every column the merge
        # covers, so styling them first is thrown away and a band comes out
        # coloured only in its first column. Styling AFTER the merge does stick
        # (the MergedCell is read-only for VALUE, not for style). Verified
        # against the written styles.xml, because openpyxl's own reader hands
        # back an unstyled MergedCell either way and cannot tell the two apart.
        f = _fill(fill)
        for c in range(c1, c2 + 1):
            tgt = ws.cell(row=r, column=c)
            tgt.border = border
            if f is not None:
                tgt.fill = f

    for i in range(ncols):
        share = grid[i + 1] - grid[i]
        ws.column_dimensions[get_column_letter(i + 1)].width = max(3.0, share * _TOTAL_WIDTH_UNITS)

    row = 1
    for b in payload.blocks:
        if b.kind == "banner":
            # the banner rides the SAME grid as the tables, or its «Date» label
            # lands in whatever hairline column happens to be second from the
            # right and gets clipped
            bb = _norm([_BANNER[0], _BANNER[1], _BANNER[2]])
            a1, z1 = _cols_for(grid, bb, 0, 1)
            a2, z2 = _cols_for(grid, bb, 1, 1)
            a3, z3 = _cols_for(grid, bb, 2, 1)
            put(row, a1, z1, f"{b.left}\n{b.leftSub}" if b.leftSub else b.left,
                bold=True, size=10, wrap=bool(b.leftSub))
            put(row, a2, z2, b.rightLabel or "Date", bold=True, align="center", fill="E5E7EB")
            put(row, a3, z3, b.rightValue, align="center")
            ws.row_dimensions[row].height = 30
            row += 1
            continue

        if b.kind in ("title", "line"):
            put(row, 1, ncols, b.text, bold=True,
                align="center" if b.kind == "title" else "left",
                size=11 if b.kind == "title" else 9)
            if b.kind == "title":
                ws.row_dimensions[row].height = 20
            row += 1
            continue

        if b.kind == "signatures":
            row += 2
            half = max(1, ncols // 2)
            for c in (1, half + 1):
                cell = ws.cell(row=row, column=c, value=_safe(b.left if c == 1 else b.right))
                cell.font = Font(name="Arial", size=9, bold=True)
            row += 3
            for c1, c2 in ((1, half), (half + 1, ncols)):
                for c in range(c1, c2 + 1):
                    tgt = ws.cell(row=row, column=c)
                    tgt.border = Border(top=thin)
            row += 2
            continue

        if b.kind != "table":
            continue

        bounds = _norm(list(b.widths)[:MAX_COLS]) if b.widths else _norm([1.0])
        for cells in b.rows[:MAX_ROWS]:
            c0 = 0
            tall = False
            last_col = 0
            for cell in cells[:MAX_COLS]:
                if c0 >= len(bounds) - 1:
                    break
                span = max(1, min(cell.span, len(bounds) - 1 - c0))
                a, z = _cols_for(grid, bounds, c0, span)
                put(row, a, z, cell.text,
                    bold=bool(cell.bold), align=cell.align or "left",
                    fill=cell.fill, wrap=bool(cell.wrap),
                    size=9 if cell.fill else 8)
                tall = tall or bool(cell.wrap and len(cell.text) > 40)
                last_col = max(last_col, z)
                c0 += span
            # A row that stopped short would leave a bare strip at the right
            # edge. Extend the LAST cell to the margin instead of parking empty
            # bordered cells there — that is what made the first attempt look
            # like it had a spare column.
            if 0 < last_col < ncols and cells:
                _extend(ws, row, last_col, ncols, border)
            if tall:
                ws.row_dimensions[row].height = 26
            row += 1
        row += 1                                   # blank line between sections

    last = get_column_letter(ncols)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0                  # may run to a second page — never squash it
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.28, right=0.28, top=0.28, bottom=0.28, header=0.1, footer=0.1)
    ws.print_area = f"A1:{last}{max(1, row)}"
    ws.freeze_panes = None

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.post("/export-excel")
async def export_sheet_excel(payload: SheetExportIn):
    try:
        data = build_sheet_workbook(payload)
    except Exception as exc:  # noqa: BLE001 — the officer needs the reason, not a 500
        raise HTTPException(status_code=422, detail=f"ساختِ فایلِ Excel ناموفق بود: {exc}")
    acc = re.sub(r"[^\w-]", "", payload.account or "") or "no-account"
    kind = "Retail" if "retail" in payload.title.lower() else "Corporate" if "corporate" in payload.title.lower() else "Summary"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="CreditFile-{kind}-{acc}.xlsx"'},
    )
